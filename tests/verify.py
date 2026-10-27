"""出力ExcelをCSVと1セルずつ突き合わせる検証スクリプト。

使い方:  python tests/verify.py [出力xlsx] [内示csv]
省略時は output/ の最新xlsx と sample/ のCSV を使う。
"""
import csv
import glob
import os
import sys
import unicodedata

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def norm(s):
    return unicodedata.normalize("NFKC", str(s or "")).replace(" ", "").upper()


def num(s):
    s = str(s or "").replace(",", "").strip()
    return float(s) if s else None


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else max(glob.glob(os.path.join(ROOT, "output", "*.xlsx")), key=os.path.getmtime)
    csv_path = sys.argv[2] if len(sys.argv) > 2 else glob.glob(os.path.join(ROOT, "sample", "*.csv"))[0]

    with open(csv_path, encoding="cp932") as f:
        rows = list(csv.reader(f))
    h = [x.strip() for x in rows[0]]
    key = h.index("整理番号")
    day_cols = {int(x): i for i, x in enumerate(h) if x.isdigit() and 1 <= int(x) <= 31}
    tot_cols = [i for i, x in enumerate(h) if x.startswith("内示数量(")][:3]
    data = {}
    for r in rows[1:]:
        if r and r[key].strip():
            data.setdefault(norm(r[key]), []).append(r)

    ws = openpyxl.load_workbook(xlsx).active
    # 見出し「背番」の位置を探す
    hdr = next(c for row in ws.iter_rows() for c in row if norm(c.value) == norm("背番"))
    hr, c0 = hdr.row, hdr.column
    headers = [ws.cell(hr, c).value for c in range(c0, ws.max_column + 1)]
    days = sum(1 for v in headers[1:] if v and "/" in str(v))
    st_col = c0 + days + 4

    ok = warn = ng = err = 0
    for r in range(hr + 1, ws.max_row + 1):
        sb = ws.cell(r, c0).value
        if sb is None:
            continue
        status = str(ws.cell(r, st_col).value or "")
        hits = data.get(norm(sb), [])
        got = [ws.cell(r, c0 + k).value for k in range(1, days + 4)]
        if not hits:
            warn += 1
            if any(v is not None for v in got) or status != "CSVに該当なし":
                err += 1
                print(f"NG 行{r} {sb}: CSVに無いのに値または判定がおかしい")
            continue

        def total(i):
            vals = [num(x[i]) for x in hits]
            return None if all(v is None for v in vals) else sum(v or 0 for v in vals)

        exp = [total(day_cols[d]) if d in day_cols else None for d in range(1, days + 1)] + [total(i) for i in tot_cols]
        got_f = [None if v is None else float(v) for v in got]
        if exp != got_f:
            err += 1
            print(f"NG 行{r} {sb}: 期待 {exp}\n           実際 {got_f}")
        if status == "一致":
            ok += 1
        elif status.startswith("日別"):
            ng += 1
        else:
            warn += 1

    print(f"一致 {ok} / 注意 {warn} / 異常 {ng} / 照合エラー {err}  ({os.path.basename(xlsx)})")
    sys.exit(1 if err else 0)


if __name__ == "__main__":
    main()
